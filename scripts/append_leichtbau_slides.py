#!/usr/bin/env python3
"""Append FiberMat CF-SMC result slides to the Leichtbaukolloquium deck."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PPTX = Path(r"C:\Users\zk0779\Documents\2026_Leichtbaukolloquium.pptx")
OUT = Path(r"C:\Users\zk0779\Documents\01_Repository\constantin_claude_FVC40\outputs")

NAVY = RGBColor(0x00, 0x2D, 0x4C)
GREEN = RGBColor(0x00, 0x96, 0x82)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x4D, 0x4D, 0x4D)
ROW = RGBColor(0xF2, 0xF2, 0xF2)

# Content band used by "Titel und Text" (layout 8)
L = Inches(0.41)
SUB_TOP = Inches(0.90)
BODY_TOP = Inches(1.32)
BODY_BOTTOM = Inches(6.85)
BODY_RIGHT = Inches(12.92)
BODY_W = BODY_RIGHT - L
BODY_H = BODY_BOTTOM - BODY_TOP


def _latin(run, name="Arial"):
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = rPr.makeelement(qn(f"a:{tag}"), {})
            rPr.append(el)
        el.set("typeface", name)


def style_run(run, *, size=18, bold=False, color=NAVY, name="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    _latin(run, name)


def set_shape_text(shape, text, *, size=18, bold=False, color=NAVY, align=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=bold, color=color)
    return tf


def add_bullets(shape, items, *, size=16, color=NAVY, space=8):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(space)
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = "•  " + item
        style_run(run, size=size, bold=False, color=color)


def drop_placeholder(shape):
    el = shape._element
    el.getparent().remove(el)


def title_subtitle(slide, title, subtitle):
    titled = subtitled = None
    extras = []
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            set_shape_text(ph, title, size=24, bold=True, color=NAVY)
            titled = ph
        elif idx in (4, 14) and subtitled is None and ph.has_text_frame:
            # subtitle band
            if ph.top < Inches(1.2):
                set_shape_text(ph, subtitle, size=16, bold=False, color=GREY)
                subtitled = ph
            else:
                extras.append(ph)
        else:
            extras.append(ph)
    return titled, subtitled, extras


def add_textbox(slide, left, top, width, height, text, *, size=16, bold=False, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets_box(slide, left, top, width, height, items, *, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    add_bullets(box, items, size=size)
    return box


def fit_picture(slide, path, left, top, width, height):
    with Image.open(path) as im:
        iw, ih = im.size
    box_a = width / height
    img_a = iw / ih
    if img_a > box_a:
        w = width
        h = int(width / img_a)
        t = top + (height - h) // 2
        l = left
    else:
        h = height
        w = int(height * img_a)
        l = left + (width - w) // 2
        t = top
    return slide.shapes.add_picture(str(path), l, t, w, h)


def shade_cell(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def write_cell(cell, text, *, size=12, bold=False, color=NAVY, fill=None, align=PP_ALIGN.CENTER):
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        shade_cell(cell, fill)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            style_run(r, size=size, bold=bold, color=color)


def add_table(slide, left, top, width, height, headers, rows):
    table = slide.shapes.add_table(1 + len(rows), len(headers), left, top, width, height).table
    for j, h in enumerate(headers):
        write_cell(table.cell(0, j), h, size=11, bold=True, color=WHITE, fill=NAVY)
    for i, row in enumerate(rows, start=1):
        bg = WHITE if i % 2 else ROW
        for j, val in enumerate(row):
            write_cell(table.cell(i, j), val, size=11, bold=False, color=NAVY, fill=bg)
    return table


def delete_from(prs, index):
    """Drop slides from ``index`` to the end (0-based)."""
    sldIdLst = prs.slides._sldIdLst
    for i in range(len(sldIdLst) - 1, index - 1, -1):
        rId = sldIdLst[i].get(qn("r:id"))
        if rId:
            prs.part.drop_rel(rId)
        sldIdLst.remove(sldIdLst[i])


def add_rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def new_content(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[8])
    titled, subtitled, extras = title_subtitle(slide, title, subtitle)
    for ph in extras:
        drop_placeholder(ph)
    add_rect(slide, L, Inches(0.82), Inches(1.4), Inches(0.055), GREEN)
    return slide


def divider(prs, title, kicker, label, image):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            set_shape_text(ph, title, size=32, bold=True, color=NAVY)
        elif idx == 14:
            set_shape_text(ph, kicker, size=14, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
        elif idx == 15:
            set_shape_text(ph, label, size=16, bold=False, color=WHITE, align=PP_ALIGN.RIGHT)
        elif idx == 13:
            ph.insert_picture(str(image))
    return slide


def main():
    prs = Presentation(str(PPTX))
    if len(prs.slides) > 33:
        first = prs.slides[33]
        title = first.shapes.title.text.strip() if first.shapes.title else ""
        if title.startswith("C-SMC"):
            delete_from(prs, 33)
    n0 = len(prs.slides)

    divider(
        prs,
        "C-SMC-Vorformlinge",
        "Anhang",
        "FiberMat",
        OUT / "preview_roll01_w18_axial.png",
    )

    s = new_content(
        prs,
        "So entstehen die Geometrien",
        "Roving ablegen  →  in Fasern teilen  →  zuschneiden  →  VTK-Datei",
    )
    add_bullets_box(
        s,
        L,
        BODY_TOP,
        Inches(6.3),
        Inches(5.3),
        [
            "Kohlenstoff-Rovings: 25,5 × 4,0 × 0,103 mm",
            "Ziel sind 40 % Faservolumen – das ist eine Obergrenze, kein Garantiewert",
            "Jeder Roving wird in 38 runde Fasern zerlegt",
            "Ausgabe sind VTK-Dateien für die DBS",
            "Sechs Rollen und drei Platten als Varianten",
        ],
        size=17,
    )
    fit_picture(
        s,
        OUT / "preview_roll01_w18_wall.png",
        Inches(6.85),
        BODY_TOP,
        Inches(6.0),
        Inches(5.2),
    )

    s = new_content(
        prs,
        "Breite und Höhe des Bündels",
        "b = 4 mm gemessen (klein)  ·  h bestimmt Subbundles und Stapelhöhe  ·  Kreuzungen fressen das FVG",
    )
    fit_picture(
        s,
        OUT / "bundle_width_height.png",
        L,
        BODY_TOP,
        BODY_W,
        Inches(3.85),
    )
    add_bullets_box(
        s,
        L,
        Inches(5.25),
        BODY_W,
        Inches(1.50),
        [
            "Breite b = 4 mm ist gemessen und klein gegen die Platte. Höhe h = 0,103 mm ist die Bündeldicke und der Faserdurchmesser.",
            "Subbundles: n = floor(b / h) = 38. Gibt man n vor, folgt h = b / n — und damit die Stapelhöhe: Deckel / h Lagen.",
            "An jeder Kreuzung sitzt die volle Höhe h über der vollen Breite 4 mm. Der Deckel ist schnell voll, die Fläche bleibt lückig → kleines FVG.",
        ],
        size=15,
    )

    s = new_content(
        prs,
        "Zylindrische C-SMC-Rollen",
        "Außen 29 mm  ·  innen 25 mm  ·  erst 30 mm breit, dann auf 18 mm geschnitten",
    )
    add_bullets_box(
        s,
        L,
        BODY_TOP,
        Inches(5.7),
        Inches(5.3),
        [
            "Wanddicke 2 mm",
            "Sechs Zufallsvarianten, gleiche Maße",
            "Die Rovings werden flach abgelegt und danach zum Ring gebogen",
            "Schnitt auf die mittleren 18 mm (von −9 bis +9 mm)",
        ],
        size=17,
    )
    fit_picture(
        s,
        OUT / "preview_roll01_w18_iso.png",
        Inches(6.3),
        BODY_TOP,
        Inches(6.5),
        Inches(5.3),
    )

    s = new_content(
        prs,
        "Rollen: Geometrie und Faservolumen",
        "Erreicht: 26 % FVG in der 2-mm-Wand  ·  Ziel: 40 %",
    )
    add_table(
        s,
        L,
        BODY_TOP,
        Inches(7.35),
        Inches(3.15),
        ["Größe", "30 mm breit", "18 mm geschnitten"],
        [
            ["Rovings", "122 – 125", "116 – 124"],
            ["Fasern", "4 636 – 4 750", "4 243 – 4 438"],
            ["Radius (mm)", "12,55 – 14,41", "12,55 – 14,41"],
            ["Höhe (mm)", "± 15", "± 9"],
            ["Soll-Ring (mm)", "innen 25 / außen 29", "gleiche Wand"],
        ],
    )
    add_bullets_box(
        s,
        L,
        Inches(4.60),
        Inches(7.35),
        Inches(2.1),
        [
            "Der Ring sitzt in der vorgegebenen Wand.",
            "In 2 mm passen nur etwa 125 Rovings. Für 40 % bräuchte man 194.",
            "Kürzere Rovings (12,5 mm) füllen etwas besser (~30 %).",
        ],
        size=15,
    )
    fit_picture(
        s,
        OUT / "preview_roll01_w18_axial.png",
        Inches(7.55),
        Inches(1.22),
        Inches(5.45),
        Inches(5.55),
    )

    s = new_content(
        prs,
        "Sechs Rollen",
        "Gleiche Maße, andere Faserlage",
    )
    fit_picture(
        s,
        OUT / "preview_six_w18.png",
        L,
        BODY_TOP,
        BODY_W,
        BODY_H,
    )

    s = new_content(
        prs,
        "C-SMC-Platten",
        "Packen in 120 × 50 × 9 mm  ·  zuschneiden auf 90 × 33 × 9 mm und 90 × 33 × 8 mm",
    )
    add_table(
        s,
        L,
        BODY_TOP,
        Inches(6.55),
        Inches(3.15),
        ["Größe", "90 × 33 × 9 mm", "90 × 33 × 8 mm"],
        [
            ["Rovings", "1 116 – 1 164", "984 – 1 027"],
            ["Fasern", "39 – 41 Tsd.", "35 – 36 Tsd."],
            ["Länge × Breite (mm)", "90 × 33", "90 × 33"],
            ["Höhe der Fasern (mm)", "bis 8,91", "bis 7,88"],
            ["FVG beim Packen", "27 %  (Ziel 40 %)", "oberstes 1 mm weg"],
        ],
    )
    add_bullets_box(
        s,
        L,
        Inches(4.60),
        Inches(6.55),
        Inches(2.1),
        [
            "Drei unabhängige Stapel. Variante 1: 1405 Rovings, 27 % FVG.",
            "Der Schnitt 90 × 33 mm sitzt genau. Die 8-mm-Platten sind oben gekürzt.",
        ],
        size=15,
    )
    fit_picture(
        s,
        OUT / "preview_cuboid_01_90x33x9.png",
        Inches(7.15),
        BODY_TOP,
        Inches(5.7),
        Inches(5.3),
    )

    s = new_content(
        prs,
        "Sechs Platten",
        "Drei Varianten  ·  zwei Dicken (9 mm und 8 mm)",
    )
    fit_picture(
        s,
        OUT / "preview_six_cuboids.png",
        L,
        BODY_TOP,
        BODY_W,
        BODY_H,
    )

    s = new_content(
        prs,
        "Warum nur etwa 27 % FVG?",
        "40 % ist nur eine Obergrenze  ·  der Stapel ist vorher schon voll",
    )
    add_table(
        s,
        L,
        BODY_TOP,
        Inches(12.4),
        Inches(2.55),
        ["", "Platte 120 × 50 × 9 mm", "Rollenwand 2 mm"],
        [
            ["Rovings für 40 %", "2 056", "194"],
            ["Rovings gelegt", "1 405", "125"],
            ["Erreichtes FVG", "27 %", "26 %"],
            ["Oberkante", "8,91 mm  (Deckel bei 9 mm)", "Wand schon voll"],
        ],
    )
    add_bullets_box(
        s,
        L,
        Inches(4.05),
        BODY_W,
        Inches(2.6),
        [
            "FiberMat legt Rovings zufällig ab. Es presst nicht.",
            "Es stoppt, wenn kein weiterer Roving unter den Deckel passt – nicht bei 40 %.",
            "Die restlichen 73 % sind Lücken zwischen den Rovings, kein leerer Raum oben.",
        ],
        size=17,
    )

    s = new_content(
        prs,
        "Warum es nicht geht",
        "Der Deckel ist schon erreicht  ·  Lücken bleiben  ·  kein weiterer Roving passt hindurch",
    )
    fit_picture(
        s,
        OUT / "fvc_why_not_slice.png",
        L,
        BODY_TOP,
        Inches(7.15),
        Inches(3.55),
    )
    fit_picture(
        s,
        OUT / "fvc_why_not_reject.png",
        Inches(7.55),
        BODY_TOP,
        Inches(5.35),
        Inches(3.55),
    )
    add_textbox(s, L, Inches(4.95), Inches(7.15), Inches(0.35),
                "Echter Schnitt  ·  rote Linie = Deckel bei 9 mm", size=13, color=GREY)
    add_bullets_box(
        s,
        L,
        Inches(5.35),
        BODY_W,
        Inches(1.40),
        [
            "Oben stößt der Stapel an den Deckel. Darunter bleiben Lücken – aber kein 25,5 × 4 mm Roving passt mehr hindurch.",
            "Deshalb stoppt das Ablegen bei 27 % FVG. Nicht weil 40 % verboten wären, sondern weil der nächste Roving nicht mehr hineinpasst.",
        ],
        size=16,
    )

    s = new_content(
        prs,
        "Zufällige Lage lässt Lücken",
        "An jeder Kreuzung kommt eine ganze Rovingdicke dazu",
    )
    fit_picture(
        s,
        OUT / "fvc_side_aligned_vs_random.png",
        L,
        BODY_TOP,
        BODY_W,
        Inches(3.55),
    )
    add_bullets_box(
        s,
        L,
        Inches(5.00),
        BODY_W,
        Inches(1.75),
        [
            "Gerade Rovings liegen in Schichten und füllen den Raum gut.",
            "Kreuzende Rovings stapeln sich. Ein langer Roving liegt oben auf und deckt die Lücken darunter zu.",
        ],
        size=16,
    )

    s = new_content(
        prs,
        "40 % FVG kommt vom Pressen, nicht vom Ablegen",
        "In einen vollen Kasten passen keine weiteren zufälligen Rovings",
    )
    fit_picture(
        s,
        OUT / "fvc_top_aligned_vs_random.png",
        L,
        BODY_TOP,
        BODY_W,
        Inches(3.35),
    )
    add_bullets_box(
        s,
        L,
        Inches(4.80),
        BODY_W,
        Inches(1.95),
        [
            "Die Restlücken sind kleiner als ein Roving (25,5 × 4 mm).",
            "Längere Rovings füllen schlechter: 12,5 mm → 30 % FVG, 25,5 mm → 27 % FVG.",
            "Dieselben Rovings in 9 mm (27 %) wären 40 %, wenn man auf 6,1 mm presst. Das macht SMC-Pressen.",
        ],
        size=16,
    )

    s = new_content(
        prs,
        "Warum FiberMat (Mahé) keine 40 % schafft",
        "Zwei Werkzeuge in einem Programm  ·  wir haben nur abgelegt, nicht gepresst",
    )
    fit_picture(
        s,
        OUT / "fvc_mahe_vs_pack.png",
        L,
        BODY_TOP,
        BODY_W,
        Inches(3.45),
    )
    add_bullets_box(
        s,
        L,
        Inches(4.90),
        BODY_W,
        Inches(1.85),
        [
            "Ursprünglich drückt FiberMat dünne Rundfasern mechanisch zusammen. Dann steigt das FVG.",
            "Bei C-SMC legt das Programm Rovings nur ab. 40 % ist eine Zielzahl, kein Pressvorgang.",
            "Der Press-Solver ist für tausende breite Rovings nicht gebaut. Dafür braucht man einen eigenen Pressschritt.",
        ],
        size=16,
    )

    s = new_content(
        prs,
        "Repo-Aufbau: Installation ok, Packen bleibt stecken",
        "7. Oktober 2026  ·  FiberMat-Zweig feat/cf-smc-stack  ·  Version 1.0.12",
    )
    add_table(
        s,
        L,
        BODY_TOP,
        Inches(12.4),
        Inches(2.85),
        ["Schritt", "Was wir gemacht haben", "Ergebnis"],
        [
            ["Klonen und Installieren", "Zweig feat/cf-smc-stack, virtuelle Umgebung", "FiberMat 1.0.12 läuft"],
            ["Test", "Kleiner Kasten 30 × 20 × 2 mm, Ziel 30 %", "60 Rovings, 30 % FVG wie gewünscht"],
            ["Sechs Rollen", "29 / 25 mm, Roving 25,5 mm, Ziel 40 %", "stecken geblieben bei 26 %"],
            ["Sechs Platten", "Kasten 120 × 50 × 9 mm, Ziel 40 %", "stecken geblieben bei 27 %"],
        ],
    )
    add_bullets_box(
        s,
        L,
        Inches(4.35),
        BODY_W,
        Inches(2.35),
        [
            "Das Programm ist installiert und hat Dateien geschrieben. Das war nicht das Problem.",
            "Wir haben Dicke (0,12 → 0,206 → 0,103 mm) und Länge (12,5 → 25,5 mm) geändert. Kürzere Rovings füllen etwas besser.",
            "40 % wurden nicht erreicht, weil der Kasten voll war – nicht weil die Installation fehlschlug.",
        ],
        size=16,
    )

    s = new_content(
        prs,
        "Kann man höheres FVG erzeugen?",
        "Ja  ·  aber nicht durch weiteres Ablegen  ·  und nicht durch einfaches Stauchen",
    )
    fit_picture(
        s,
        OUT / "fvc_compact_options.png",
        L,
        BODY_TOP,
        BODY_W,
        Inches(3.45),
    )
    add_bullets_box(
        s,
        L,
        Inches(4.90),
        BODY_W,
        Inches(1.85),
        [
            "pack() presst nicht. Nur in z stauchen lässt Fasern sich durchdringen – das darf die DBS nicht als Start bekommen.",
            "Echtes Pressen: der Solver biegt Rovings in die Lücken. Bei tausenden C-SMC-Rovings ist das zu schwer.",
            "26–27 % ist der sinnvolle Start in der Presse. 40 % ist der Zustand nach dem Pressen.",
        ],
        size=15,
    )

    s = new_content(
        prs,
        "Fazit für die DBS",
        "Die Geometrie ist nutzbar  ·  das Faservolumen kommt vom Ablegen, nicht von den 40 %",
    )
    add_bullets_box(
        s,
        L,
        BODY_TOP,
        BODY_W,
        Inches(5.4),
        [
            "Sechs Rollen (außen 29 mm, innen 25 mm) sitzen in der Soll-Wand. Schnitt auf 18 mm genau.",
            "Sechs Platten: 90 × 33 × 9 mm und 90 × 33 × 8 mm.",
            "Zufälliges Ablegen bleibt bei 26–27 % FVG. 40 % nur durch echtes Pressen oder Ausrichten, nicht durch Stauchen.",
            "Jeder 4-mm-Roving wird 38 runde Fasern. Farbe = Faserwinkel.",
            "Die Varianten unterscheiden sich in der Faserlage, nicht in den Maßen – geeignet als DBS-Satz.",
        ],
        size=17,
    )

    try:
        prs.save(str(PPTX))
        saved = PPTX
    except PermissionError:
        saved = PPTX.with_name(PPTX.stem + "_updated.pptx")
        prs.save(str(saved))
    print(f"appended {len(prs.slides) - n0} slides  total {len(prs.slides)}  saved {saved}")


if __name__ == "__main__":
    main()
